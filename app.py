from __future__ import annotations

import argparse
import cgi
import html
import io
import json
import os
import tempfile
import uuid
from dataclasses import dataclass
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

import easyocr
import numpy as np
import pypdfium2 as pdfium
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image

HOST = "0.0.0.0"
PORT = 5000
SUPPORTED_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".pdf"}
EXCEL_CACHE: dict[str, tuple[str, bytes]] = {}


@dataclass
class OCRCell:
    text: str
    start_row: int
    end_row: int
    start_col: int
    end_col: int
    confidence: float


@dataclass
class OCRSheetLayout:
    title: str
    row_count: int
    col_count: int
    cells: list[OCRCell]


class HandwritingToExcelConverter:
    def __init__(self, languages: list[str], gpu: bool = False) -> None:
        self.reader = easyocr.Reader(languages, gpu=gpu)

    def process_document(self, input_path: Path, confidence_threshold: float = 0.2) -> list[OCRSheetLayout]:
        pages = self._load_pages(input_path)
        layouts: list[OCRSheetLayout] = []

        for index, image in enumerate(pages, start=1):
            title = f"Page_{index}"
            layouts.append(self._analyze_page(image, title, confidence_threshold))

        return layouts

    def create_workbook_bytes(self, layouts: list[OCRSheetLayout]) -> bytes:
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        for layout in layouts:
            sheet = workbook.create_sheet(layout.title)
            self._write_layout_to_sheet(sheet, layout)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _load_pages(self, input_path: Path) -> list[Image.Image]:
        ext = input_path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"不支援的檔案格式：{ext}")

        if ext == ".pdf":
            return self._render_pdf_pages(input_path)

        return [Image.open(input_path).convert("RGB")]

    def _render_pdf_pages(self, pdf_path: Path) -> list[Image.Image]:
        pdf = pdfium.PdfDocument(str(pdf_path))
        pages: list[Image.Image] = []

        for index in range(len(pdf)):
            page = pdf.get_page(index)
            bitmap = page.render(scale=220 / 72)
            pages.append(bitmap.to_pil().convert("RGB"))
            page.close()

        pdf.close()
        return pages

    def _analyze_page(self, image: Image.Image, title: str, confidence_threshold: float) -> OCRSheetLayout:
        image_array = np.array(image)
        height, width = image_array.shape[:2]

        results = self.reader.readtext(
            image_array,
            detail=1,
            paragraph=False,
            text_threshold=0.5,
            low_text=0.2,
            width_ths=0.7,
        )

        filtered = [item for item in results if item[2] >= confidence_threshold and item[1].strip()]

        row_count = min(max(int(height / 45), 20), 220)
        col_count = min(max(int(width / 90), 12), 120)
        row_height_px = height / row_count
        col_width_px = width / col_count

        cells: list[OCRCell] = []
        for bbox, text, confidence in filtered:
            xs = [point[0] for point in bbox]
            ys = [point[1] for point in bbox]
            x_min, x_max = max(min(xs), 0), min(max(xs), width)
            y_min, y_max = max(min(ys), 0), min(max(ys), height)

            start_col = max(1, min(col_count, int(x_min / col_width_px) + 1))
            end_col = max(1, min(col_count, int(np.ceil(x_max / col_width_px))))
            start_row = max(1, min(row_count, int(y_min / row_height_px) + 1))
            end_row = max(1, min(row_count, int(np.ceil(y_max / row_height_px))))

            cells.append(
                OCRCell(
                    text=text,
                    start_row=start_row,
                    end_row=end_row,
                    start_col=start_col,
                    end_col=end_col,
                    confidence=float(confidence),
                )
            )

        return OCRSheetLayout(title=title, row_count=row_count, col_count=col_count, cells=cells)

    def _write_layout_to_sheet(self, sheet, layout: OCRSheetLayout) -> None:
        for row in range(1, layout.row_count + 1):
            sheet.row_dimensions[row].height = 18

        for col in range(1, layout.col_count + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 12

        for cell_data in layout.cells:
            cell = sheet.cell(row=cell_data.start_row, column=cell_data.start_col, value=cell_data.text)
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.font = Font(size=11)
            if cell_data.confidence < 0.4:
                cell.font = Font(size=11, color="FF7F50")

            if cell_data.end_row > cell_data.start_row or cell_data.end_col > cell_data.start_col:
                try:
                    sheet.merge_cells(
                        start_row=cell_data.start_row,
                        start_column=cell_data.start_col,
                        end_row=cell_data.end_row,
                        end_column=cell_data.end_col,
                    )
                except ValueError:
                    continue


def build_preview_html(layouts: list[OCRSheetLayout]) -> str:
    pages_html = [render_sheet_html(layout) for layout in layouts]
    return "\n".join(pages_html)


def render_sheet_html(layout: OCRSheetLayout) -> str:
    grid: dict[tuple[int, int], OCRCell] = {}
    covered: set[tuple[int, int]] = set()

    for cell in layout.cells:
        grid[(cell.start_row, cell.start_col)] = cell
        for row in range(cell.start_row, cell.end_row + 1):
            for col in range(cell.start_col, cell.end_col + 1):
                if row == cell.start_row and col == cell.start_col:
                    continue
                covered.add((row, col))

    rows: list[str] = []
    max_preview_rows = min(layout.row_count, 60)
    max_preview_cols = min(layout.col_count, 32)

    for row in range(1, max_preview_rows + 1):
        tds: list[str] = []
        for col in range(1, max_preview_cols + 1):
            if (row, col) in covered:
                continue

            cell = grid.get((row, col))
            if cell is None:
                tds.append('<td class="empty"></td>')
                continue

            rowspan = min(cell.end_row, max_preview_rows) - cell.start_row + 1
            colspan = min(cell.end_col, max_preview_cols) - cell.start_col + 1
            safe_text = html.escape(cell.text).replace("\n", "<br>")
            cls = "low-confidence" if cell.confidence < 0.4 else ""
            tds.append(
                f'<td class="{cls}" rowspan="{rowspan}" colspan="{colspan}" title="信心值 {cell.confidence:.2f}">{safe_text}</td>'
            )

        rows.append(f"<tr>{''.join(tds)}</tr>")

    table_rows = "".join(rows)
    return (
        f'<section class="sheet">'
        f'<h3>{html.escape(layout.title)}</h3>'
        f'<div class="sheet-wrap"><table class="excel-like">{table_rows}</table></div>'
        f"</section>"
    )


class OCRWebHandler(BaseHTTPRequestHandler):
    converter = HandwritingToExcelConverter(languages=["ch_tra", "en"], gpu=False)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/":
            self._send_html(HTTPStatus.OK, self._render_index())
            return

        if parsed.path == "/download":
            self._handle_download(parsed)
            return

        self.send_error(HTTPStatus.NOT_FOUND, "Not Found")

    def do_POST(self) -> None:
        if self.path != "/analyze":
            self.send_error(HTTPStatus.NOT_FOUND, "Not Found")
            return

        try:
            payload = self._handle_upload()
            self._send_json(HTTPStatus.OK, payload)
        except ValueError as exc:
            self._send_json(HTTPStatus.BAD_REQUEST, {"error": str(exc)})
        except Exception as exc:  # noqa: BLE001
            self._send_json(HTTPStatus.INTERNAL_SERVER_ERROR, {"error": f"處理失敗：{exc}"})

    def _handle_upload(self) -> dict[str, Any]:
        content_type = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in content_type:
            raise ValueError("請使用表單上傳檔案。")

        form = cgi.FieldStorage(
            fp=self.rfile,
            headers=self.headers,
            environ={
                "REQUEST_METHOD": "POST",
                "CONTENT_TYPE": content_type,
            },
        )

        file_item = form["file"] if "file" in form else None
        if file_item is None or not getattr(file_item, "filename", ""):
            raise ValueError("請上傳圖片或 PDF 檔案。")

        confidence = 0.2
        if "confidence_threshold" in form:
            confidence = float(form.getvalue("confidence_threshold"))

        filename = Path(file_item.filename).name
        ext = Path(filename).suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"不支援的副檔名：{ext}")

        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
            temp_file.write(file_item.file.read())
            temp_path = Path(temp_file.name)

        try:
            layouts = self.converter.process_document(temp_path, confidence_threshold=confidence)
            preview_html = build_preview_html(layouts)
            excel_bytes = self.converter.create_workbook_bytes(layouts)
        finally:
            if temp_path.exists():
                os.remove(temp_path)

        token = uuid.uuid4().hex
        excel_name = f"{Path(filename).stem}_ocr.xlsx"
        EXCEL_CACHE[token] = (excel_name, excel_bytes)

        return {
            "preview_html": preview_html,
            "download_url": f"/download?token={token}",
            "excel_filename": excel_name,
            "sheet_count": len(layouts),
        }

    def _handle_download(self, parsed) -> None:
        query = parse_qs(parsed.query)
        token = query.get("token", [""])[0]
        if not token or token not in EXCEL_CACHE:
            self.send_error(HTTPStatus.NOT_FOUND, "File not found")
            return

        filename, content = EXCEL_CACHE[token]
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Length", str(len(content)))
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.end_headers()
        self.wfile.write(content)

    def _send_html(self, status: HTTPStatus, body: str) -> None:
        payload = body.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def _send_json(self, status: HTTPStatus, payload: dict[str, Any]) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _render_index(self) -> str:
        return """<!doctype html>
<html lang=\"zh-Hant\">
<head>
  <meta charset=\"UTF-8\" />
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\" />
  <title>中文手寫 OCR 轉 Excel</title>
  <style>
    body { font-family: Arial, sans-serif; margin: 0; background: #f5f7fb; color: #1f2937; }
    .container { max-width: 1200px; margin: 24px auto; padding: 24px; background: #fff; border-radius: 12px; box-shadow: 0 8px 24px rgba(15, 23, 42, 0.12); }
    h1 { margin-top: 0; }
    .dropzone { border: 2px dashed #60a5fa; border-radius: 10px; background: #eff6ff; padding: 30px; text-align: center; cursor: pointer; }
    .dropzone.dragover { border-color: #1d4ed8; background: #dbeafe; }
    .controls { margin-top: 12px; display: flex; gap: 16px; align-items: center; flex-wrap: wrap; }
    .btn { margin-top: 14px; background: #2563eb; color: #fff; border: 0; border-radius: 8px; padding: 10px 14px; cursor: pointer; }
    .btn:disabled { opacity: .6; cursor: not-allowed; }
    #status { margin-top: 12px; font-weight: 700; }
    #preview { margin-top: 20px; }
    .sheet { margin-top: 24px; }
    .sheet-wrap { border: 1px solid #cbd5e1; overflow: auto; max-height: 620px; background: white; }
    .excel-like { border-collapse: collapse; table-layout: fixed; }
    .excel-like td { min-width: 90px; height: 28px; border: 1px solid #e2e8f0; padding: 4px 6px; vertical-align: top; white-space: pre-wrap; }
    .excel-like td.empty { background: #f8fafc; }
    .low-confidence { color: #c2410c; }
    .download-link { display: inline-block; margin-top: 12px; font-weight: 700; color: #1d4ed8; }
  </style>
</head>
<body>
  <div class=\"container\">
    <h1>中文手寫字辨識（圖片/PDF → Excel）</h1>
    <p>拖曳檔案進來後，會在下方先顯示近似 Excel 的內容預覽，同時可下載 `.xlsx`。</p>

    <div id=\"dropzone\" class=\"dropzone\">
      <p><strong>拖曳圖片或 PDF 到這裡，或點擊選擇檔案</strong></p>
      <p>支援：png/jpg/jpeg/bmp/tif/tiff/pdf</p>
      <input id=\"fileInput\" type=\"file\" accept=\".png,.jpg,.jpeg,.bmp,.tif,.tiff,.pdf\" hidden />
      <p id=\"selectedFile\">尚未選擇檔案</p>
    </div>

    <div class=\"controls\">
      <label>信心值門檻：
        <input id=\"confidence\" type=\"number\" min=\"0\" max=\"1\" step=\"0.05\" value=\"0.2\" />
      </label>
    </div>

    <button id=\"analyzeBtn\" class=\"btn\" disabled>開始辨識</button>
    <div id=\"status\"></div>
    <a id=\"downloadLink\" class=\"download-link\" href=\"#\" style=\"display:none\">下載 Excel</a>

    <div id=\"preview\"></div>
  </div>

  <script>
    const dropzone = document.getElementById('dropzone');
    const fileInput = document.getElementById('fileInput');
    const analyzeBtn = document.getElementById('analyzeBtn');
    const selectedFile = document.getElementById('selectedFile');
    const statusEl = document.getElementById('status');
    const previewEl = document.getElementById('preview');
    const confidenceEl = document.getElementById('confidence');
    const downloadLink = document.getElementById('downloadLink');

    let chosenFile = null;

    function setFile(file) {
      chosenFile = file;
      if (file) {
        selectedFile.textContent = `已選擇：${file.name}`;
        analyzeBtn.disabled = false;
      } else {
        selectedFile.textContent = '尚未選擇檔案';
        analyzeBtn.disabled = true;
      }
    }

    dropzone.addEventListener('click', () => fileInput.click());
    fileInput.addEventListener('change', () => setFile(fileInput.files[0] || null));

    ['dragenter', 'dragover'].forEach(eventName => {
      dropzone.addEventListener(eventName, (e) => {
        e.preventDefault();
        dropzone.classList.add('dragover');
      });
    });

    ['dragleave', 'drop'].forEach(eventName => {
      dropzone.addEventListener(eventName, (e) => {
        e.preventDefault();
        dropzone.classList.remove('dragover');
      });
    });

    dropzone.addEventListener('drop', (e) => {
      const files = e.dataTransfer.files;
      if (!files || files.length === 0) return;
      setFile(files[0]);
    });

    analyzeBtn.addEventListener('click', async () => {
      if (!chosenFile) return;

      statusEl.textContent = '辨識中，請稍候...';
      previewEl.innerHTML = '';
      downloadLink.style.display = 'none';
      analyzeBtn.disabled = true;

      const form = new FormData();
      form.append('file', chosenFile);
      form.append('confidence_threshold', confidenceEl.value || '0.2');

      try {
        const resp = await fetch('/analyze', { method: 'POST', body: form });
        const data = await resp.json();

        if (!resp.ok) {
          statusEl.textContent = data.error || '辨識失敗';
          return;
        }

        statusEl.textContent = `完成：共 ${data.sheet_count} 頁。下方為 Excel 版面預覽。`;
        previewEl.innerHTML = data.preview_html;
        downloadLink.href = data.download_url;
        downloadLink.textContent = `下載 ${data.excel_filename}`;
        downloadLink.style.display = 'inline-block';
      } catch (err) {
        statusEl.textContent = `發生錯誤：${err}`;
      } finally {
        analyzeBtn.disabled = false;
      }
    });
  </script>
</body>
</html>"""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="中文手寫字辨識 Web UI")
    parser.add_argument("--host", default=HOST, help="服務監聽主機")
    parser.add_argument("--port", type=int, default=PORT, help="服務埠號")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    server = ThreadingHTTPServer((args.host, args.port), OCRWebHandler)
    print(f"服務已啟動：http://{args.host}:{args.port}")
    server.serve_forever()


if __name__ == "__main__":
    main()
